from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import rasterio
from rasterio.warp import transform_bounds
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data" / "raw"
DOCS_DIR = ROOT / "docs"
OUTPUT = DOCS_DIR / "data_inventory.csv"


HEADERS = [
    "图层类型",
    "文件路径",
    "文件名",
    "年份",
    "格式",
    "投影_CRS",
    "空间范围_WGS84",
    "分辨率",
    "波段数",
    "数据类型",
    "nodata",
    "行列或尺寸",
    "表格结构",
    "建议用途",
    "是否需要预处理",
    "预处理建议",
    "备注",
]


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT)).replace("\\", "/")


def years_from_name(name: str) -> str:
    years = sorted(set(re.findall(r"(?:19|20)\d{2}", name)))
    return ";".join(years)


def layer_type(path: Path) -> str:
    parent = path.parent.name
    name = path.name
    text = f"{parent}/{name}"
    if "土地利用" in text or "CLCD" in text:
        return "土地利用_CLCD_30m"
    if "夜间灯光" in text or "LongNTL" in text:
        return "夜间灯光_LongNTL_500m"
    if "DEM" in text:
        return "DEM_NASA_30m"
    if "人口" in text or "LandScan" in text:
        return "人口_LandScan_1km"
    if "GDP" in text:
        return "GDP_1km"
    if "NDVI" in text:
        return "NDVI_MODIS_250m"
    if "EVI" in text:
        return "EVI_MODIS_250m"
    if "水体" in text or "JRC" in text:
        return "水体_JRC_30m"
    if "TerraClimate" in text:
        return "TerraClimate_气候水文"
    if "气温降水" in text or path.suffix.lower() == ".xlsx":
        return "气象站点_气温降水"
    if "boundary" in text or path.suffix.lower() in {".geojson", ".json"}:
        return "研究区边界"
    if path.suffix.lower() == ".pdf":
        return "项目文档"
    if path.suffix.lower() == ".txt":
        return "数据说明文档"
    return "其他"


def suggested_use(layer: str) -> str:
    uses = {
        "土地利用_CLCD_30m": "土地利用变化、生态用地比例、景观格局、转移矩阵、抵抗力指标",
        "夜间灯光_LongNTL_500m": "人类活动强度、城镇化水平、开发压力、适应力指标",
        "DEM_NASA_30m": "三维地形底座、坡度坡向、地形约束、洪水暴露分析",
        "人口_LandScan_1km": "人口暴露度、人口密度、社会经济适应力",
        "GDP_1km": "经济发展水平、社会经济适应力、区域发展差异",
        "NDVI_MODIS_250m": "植被覆盖、生态质量、稳定性、恢复力趋势",
        "EVI_MODIS_250m": "植被活力、生态恢复、生态质量评价",
        "水体_JRC_30m": "水体分布、水体频率、洪泛暴露、水域变化",
        "TerraClimate_气候水文": "气候水文背景、水分亏缺、干湿状况、韧性气候因子",
        "气象站点_气温降水": "站点气候背景、栅格气候验证、降水气温对照",
        "研究区边界": "地图高亮、裁剪、分区统计、PostGIS 入库",
        "项目文档": "技术路线、系统架构、里程碑参考",
        "数据说明文档": "数据来源、年份、分辨率、使用说明参考",
    }
    return uses.get(layer, "")


def preprocessing(layer: str, suffix: str) -> tuple[str, str]:
    if suffix.lower() == ".tif":
        if "土地利用" in layer or "水体" in layer:
            return "是", "统一投影、范围和目标网格；分类数据重采样使用最近邻；必要时转换为 COG"
        return "是", "统一投影、范围和目标网格；连续数据按目标分辨率重采样；必要时标准化并转换为 COG"
    if suffix.lower() in {".csv", ".xlsx"}:
        return "是", "统一字段名、年份字段、坐标字段和编码；必要时导入 PostGIS 或转换为分析表"
    if suffix.lower() in {".geojson", ".json"}:
        return "视情况", "检查几何有效性和坐标系；必要时简化边界并建立空间索引"
    return "否", ""


def raster_row(path: Path) -> dict[str, str]:
    layer = layer_type(path)
    need, prep = preprocessing(layer, path.suffix)
    with rasterio.open(path) as src:
        crs = str(src.crs) if src.crs else ""
        bounds = ""
        if src.crs:
            wb = transform_bounds(src.crs, "EPSG:4326", *src.bounds, densify_pts=21)
            bounds = f"{wb[0]:.6f},{wb[1]:.6f},{wb[2]:.6f},{wb[3]:.6f}"
        res = ";".join(f"{x:.12g}" for x in src.res)
        dtype = ";".join(src.dtypes)
        nodata = "" if src.nodata is None else str(src.nodata)
        size = f"{src.width}x{src.height}"
        bands = str(src.count)
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": "GeoTIFF",
        "投影_CRS": crs,
        "空间范围_WGS84": bounds,
        "分辨率": res,
        "波段数": bands,
        "数据类型": dtype,
        "nodata": nodata,
        "行列或尺寸": size,
        "表格结构": "",
        "建议用途": suggested_use(layer),
        "是否需要预处理": need,
        "预处理建议": prep,
        "备注": "",
    }


def csv_row(path: Path) -> dict[str, str]:
    layer = layer_type(path)
    need, prep = preprocessing(layer, path.suffix)
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, [])
        rows = sum(1 for _ in reader)
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": "CSV",
        "投影_CRS": "",
        "空间范围_WGS84": "",
        "分辨率": "",
        "波段数": "",
        "数据类型": "table",
        "nodata": "",
        "行列或尺寸": f"{rows} rows x {len(header)} cols",
        "表格结构": ";".join(header),
        "建议用途": suggested_use(layer),
        "是否需要预处理": need,
        "预处理建议": prep,
        "备注": "",
    }


def xlsx_row(path: Path) -> dict[str, str]:
    layer = layer_type(path)
    need, prep = preprocessing(layer, path.suffix)
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"{sheet_name}:{ws.max_row}x{ws.max_column}")
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": "Excel",
        "投影_CRS": "",
        "空间范围_WGS84": "",
        "分辨率": "",
        "波段数": "",
        "数据类型": "table",
        "nodata": "",
        "行列或尺寸": ";".join(parts),
        "表格结构": "工作表:" + ";".join(wb.sheetnames),
        "建议用途": suggested_use(layer),
        "是否需要预处理": need,
        "预处理建议": prep,
        "备注": "",
    }


def geojson_row(path: Path) -> dict[str, str]:
    layer = layer_type(path)
    need, prep = preprocessing(layer, path.suffix)
    data = json.loads(path.read_text(encoding="utf-8"))
    features = data.get("features", [])
    bbox = data.get("bbox", [])
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": "GeoJSON",
        "投影_CRS": "CRS84/WGS84",
        "空间范围_WGS84": ",".join(str(x) for x in bbox),
        "分辨率": "",
        "波段数": "",
        "数据类型": "vector",
        "nodata": "",
        "行列或尺寸": f"{len(features)} features",
        "表格结构": "",
        "建议用途": suggested_use(layer),
        "是否需要预处理": need,
        "预处理建议": prep,
        "备注": "由裁剪栅格有效区域反推生成",
    }


def pdf_row(path: Path) -> dict[str, str]:
    reader = PdfReader(str(path))
    layer = layer_type(path)
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": "PDF",
        "投影_CRS": "",
        "空间范围_WGS84": "",
        "分辨率": "",
        "波段数": "",
        "数据类型": "document",
        "nodata": "",
        "行列或尺寸": f"{len(reader.pages)} pages",
        "表格结构": "",
        "建议用途": suggested_use(layer),
        "是否需要预处理": "否",
        "预处理建议": "",
        "备注": "",
    }


def text_row(path: Path) -> dict[str, str]:
    layer = layer_type(path)
    return {
        "图层类型": layer,
        "文件路径": rel(path),
        "文件名": path.name,
        "年份": years_from_name(path.name),
        "格式": path.suffix.lstrip(".").upper(),
        "投影_CRS": "",
        "空间范围_WGS84": "",
        "分辨率": "",
        "波段数": "",
        "数据类型": "document",
        "nodata": "",
        "行列或尺寸": f"{path.stat().st_size} bytes",
        "表格结构": "",
        "建议用途": suggested_use(layer),
        "是否需要预处理": "否",
        "预处理建议": "",
        "备注": "",
    }


def build_rows() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    paths = [
        p
        for p in sorted(RAW_DIR.rglob("*"))
        if p.is_file() and p.name not in {".DS_Store"} and p.suffix.lower() != ".gitkeep"
    ]
    paths.extend(sorted(DOCS_DIR.glob("*.pdf")))

    for path in paths:
        suffix = path.suffix.lower()
        if suffix == ".tif":
            rows.append(raster_row(path))
        elif suffix == ".csv":
            rows.append(csv_row(path))
        elif suffix == ".xlsx":
            rows.append(xlsx_row(path))
        elif suffix in {".geojson", ".json"}:
            rows.append(geojson_row(path))
        elif suffix == ".pdf":
            rows.append(pdf_row(path))
        elif suffix == ".txt":
            rows.append(text_row(path))
    return rows


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    rows = build_rows()
    with OUTPUT.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {OUTPUT}")
    print(f"rows {len(rows)}")


if __name__ == "__main__":
    main()
