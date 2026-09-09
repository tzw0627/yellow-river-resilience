from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from ..context import layer_label, load_summary, timeline_stat

router = APIRouter(prefix="/api/report", tags=["report"])


class ReportRequest(BaseModel):
    year: int
    layers: list[str]
    title: str = "黄河滩区生态韧性分析报告"
    format: str = "docx"  # docx | pdf


def _collect_rows(year: int, layers: list[str]) -> list[tuple[str, dict]]:
    rows: list[tuple[str, dict]] = []
    for layer in layers:
        stat = timeline_stat(year, layer)
        if stat:
            rows.append((layer, stat))
    return rows


def _fmt(value) -> str:
    if value is None:
        return "--"
    try:
        return f"{float(value):.4f}"
    except (TypeError, ValueError):
        return str(value)


def _build_docx(req: ReportRequest, rows: list[tuple[str, dict]]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(req.title, level=0)
    doc.add_paragraph(f"研究区：黄河中下游（河南段）滩区")
    doc.add_paragraph(f"分析年份：{req.year}")
    doc.add_paragraph(f"生成时间：{datetime.now():%Y-%m-%d %H:%M}")

    bounds = load_summary().get("bounds", {})
    doc.add_paragraph(
        f"范围：经度 {bounds.get('west')}~{bounds.get('east')}，纬度 {bounds.get('south')}~{bounds.get('north')}。"
    )

    doc.add_heading("图层统计概览", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, name in enumerate(["图层", "中位数", "最小值", "最大值", "P90", "有效像元比例"]):
        header[i].text = name
        header[i].paragraphs[0].runs[0].font.size = Pt(10)

    for layer, stat in rows:
        cells = table.add_row().cells
        cells[0].text = f"{layer_label(layer)}（{layer}）"
        cells[1].text = _fmt(stat.get("median"))
        cells[2].text = _fmt(stat.get("min"))
        cells[3].text = _fmt(stat.get("max"))
        cells[4].text = _fmt(stat.get("p90"))
        ratio = stat.get("validRatio") or stat.get("valid_ratio")
        cells[5].text = f"{float(ratio) * 100:.1f}%" if ratio is not None else "--"

    doc.add_heading("说明", level=1)
    doc.add_paragraph(
        "本报告由平台后端自动生成，统计值来自统一 250m / EPSG:4326 网格的真实数据产品。"
        "四维生态韧性（ER/ERS/ERD/ERM/ERF）数值越高表示韧性越强，洪水风险 FRI 数值越高表示风险越高。"
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_pdf(req: ReportRequest, rows: list[tuple[str, dict]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = getSampleStyleSheet()
    title_style = ParagraphStyle("CNTitle", parent=base["Title"], fontName="STSong-Light")
    h_style = ParagraphStyle("CNHeading", parent=base["Heading2"], fontName="STSong-Light")
    body = ParagraphStyle("CNBody", parent=base["Normal"], fontName="STSong-Light", fontSize=10, leading=16)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=req.title)
    bounds = load_summary().get("bounds", {})
    story = [
        Paragraph(req.title, title_style),
        Spacer(1, 8),
        Paragraph(f"研究区：黄河中下游（河南段）滩区", body),
        Paragraph(f"分析年份：{req.year}", body),
        Paragraph(f"生成时间：{datetime.now():%Y-%m-%d %H:%M}", body),
        Paragraph(
            f"范围：经度 {bounds.get('west')}~{bounds.get('east')}，纬度 {bounds.get('south')}~{bounds.get('north')}。",
            body,
        ),
        Spacer(1, 12),
        Paragraph("图层统计概览", h_style),
    ]

    data = [["图层", "中位数", "最小值", "最大值", "P90", "有效像元比例"]]
    for layer, stat in rows:
        ratio = stat.get("validRatio") or stat.get("valid_ratio")
        data.append(
            [
                f"{layer_label(layer)}",
                _fmt(stat.get("median")),
                _fmt(stat.get("min")),
                _fmt(stat.get("max")),
                _fmt(stat.get("p90")),
                f"{float(ratio) * 100:.1f}%" if ratio is not None else "--",
            ]
        )

    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d6f88")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d8d0c0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f1e8")]),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            "本报告由平台后端自动生成，统计值来自统一 250m / EPSG:4326 网格的真实数据产品。",
            body,
        )
    )
    doc.build(story)
    return buffer.getvalue()


@router.post("")
def generate_report(req: ReportRequest):
    rows = _collect_rows(req.year, req.layers)
    if not rows:
        raise HTTPException(status_code=400, detail="所选图层在该年份没有可用统计")

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if req.format == "pdf":
        content = _build_pdf(req, rows)
        media = "application/pdf"
        filename = f"resilience_report_{req.year}_{stamp}.pdf"
    else:
        content = _build_docx(req, rows)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"resilience_report_{req.year}_{stamp}.docx"

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(io.BytesIO(content), media_type=media, headers=headers)
